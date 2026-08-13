"""
Build a dispatch sheet (CSV + XLSX) and an interactive 3D-loading / route-map HTML report
from one day's order-line CSV (as produced by extract_orders.ps1).

Usage:
    python build_dispatch.py --csv output\\orders_2026-07-08.csv --date 2026-07-08 ^
        --template lib\\template.html --out-html output\\dispatch_2026-07-08.html ^
        --out-csv output\\dispatch_sheet_2026-07-08.csv --out-xlsx output\\dispatch_sheet_2026-07-08.xlsx ^
        --truck-l 8700 --truck-w 2400 --truck-h 2400

run.ps1 wraps this call -- you normally don't need to invoke it directly.
"""
import argparse
import csv
import json
import math
import os
import sys
import urllib.request
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Some customer names use old 自造字 (private-use-area Unicode characters not in the console's
# codepage) -- without this, printing a progress line for that one customer crashes the whole
# run instead of just showing a "?" in place of the one character.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(errors="replace")
    except AttributeError:
        pass


# pallet mode: standard pallet SKUs a box gets snapped to (mm), base height reserved under
# the cargo, and the shared 2200mm total-height cap both modes stack up against (matches the
# demo Artifact's PALLET_OPTIONS/PALLET_BASE_HEIGHT/PALLET_MAX_TOTAL_HEIGHT exactly, so the
# same real-world truck load produces the same numbers in both tools).
PALLET_OPTIONS = [(1100.0, 1100.0), (1100.0, 1300.0)]
PALLET_BASE_HEIGHT = 150.0
MAX_TOTAL_HEIGHT = 2200.0

# plant code -> display name, matching config.json's db.database value and the demo Artifact's
# own PLANT_INFO table (which also carries each plant's depot lat/lng for route distance) --
# the HTML report's client-side engine already has all four plants built in, so a report only
# ever needs to say WHICH one this data belongs to, not repeat the address/coordinates here.
PLANT_NAMES = {"T2": "龍潭廠", "T10": "雲林廠", "T3": "神岡廠", "T6": "路竹廠"}


def choose_pallet(w, l, overhang_w, overhang_l):
    """Smallest standard pallet SKU (from PALLET_OPTIONS) the box fits on, allowing it to
    overhang the pallet edge by up to overhang_w/overhang_l per axis. None if the box is too
    big for any of them (falls back to using the box's own footprint as its "pallet")."""
    for pw, pl in PALLET_OPTIONS:
        eff_pw, eff_pl = pw + overhang_w, pl + overhang_l
        if (w <= eff_pw and l <= eff_pl) or (w <= eff_pl and l <= eff_pw):
            return pw, pl
    return None


def columns_per_pallet(bw, bl, eff_pw, eff_pl):
    """How many boxes fit side by side on one pallet layer, trying both box orientations."""
    opt1 = math.floor(eff_pw / bw) * math.floor(eff_pl / bl)
    opt2 = math.floor(eff_pw / bl) * math.floor(eff_pl / bw)
    return max(1, opt1, opt2)


def pack_manual_columns(order_lines, max_total_height, mode="greedy"):
    """手工疊車 only: 同客戶相鄰裝車序號可以疊在另一筆訂單的上方,只要不超過 max_total_height。
    Fills height budget across ORDERS at the box-quantity level, not as an all-or-nothing merge
    of whole pre-sized stacks: a customer's later, smaller-footprint order first contributes as
    many of its OWN boxes as fit into whatever headroom existing columns have left (riding on
    top of them), and only whatever quantity doesn't fit that way starts fresh column(s) of its
    own at the full height budget. E.g. order A fills a column to 1600/2200mm; order B (smaller
    footprint, larger quantity) puts as many boxes as fit in the remaining 600mm on top of A,
    and the rest of B's boxes become their own separate column -- so one order can end up split
    across a shared column AND a dedicated one, which a whole-stack-at-a-time merge can't do.

    Different customers never share a column (nobody wants a stranger's freight resting on
    theirs, and it breaks sequential unloading). A column's footprint is fixed by whatever
    order started it; a rider just needs to fit within that footprint, checked via sorted
    dimension pairs so it doesn't matter which physical orientation the packer later places
    the base in.

    Runs as a pre-placement pass (grouped by customer, decided before any truck is chosen)
    rather than opportunistically during placement the way the demo Artifact's JS does it --
    this tool's packer is a flat 2D shelf packer with no "column" concept to hook into, so this
    achieves the same physical/business outcome without rebuilding the packer around a
    stacking-aware placement pass. A column base's own "height" grows to the full column
    height (so the existing volume/placement math needs no changes at all); its true own
    height is stashed in "_own_height" and each rider is stashed in "_children" (with its own
    y0 offset) so the report-building step can split them back into separate rows later.

    `order_lines` are UNCHUNKED: one dict per order line (order_no, cust_code, cust_name, lat,
    lng, route, main_road, item_code, box_w, box_l, t, qty) -- chunking into height-capped
    stacks happens here, per column, instead of upfront, since how much of an order's quantity
    goes on an existing column vs. a fresh one isn't known until this pass runs.

    `mode`: "greedy" (default) is the ordinary same-customer on-top-stacking behaviour above.
    "none" skips riding entirely -- every line becomes its own independent column. Used by
    build()'s joint stacking/floor-placement comparison: a customer whose orders happen to tile
    the floor well side by side can come out ahead with NO height-stacking at all, which greedy
    mode has no way to discover since it commits to a stacking decision before floor placement
    (try_place) ever runs.
    """
    by_cust = {}
    for ol in order_lines:
        by_cust.setdefault(ol["cust_code"], []).append(ol)

    stacks = []
    for cust_lines in by_cust.values():
        # largest footprint first -- a rider can only go on top of something whose footprint
        # it fits WITHIN, so the biggest orders have to become column bases before anything
        # smaller is considered for them.
        cust_lines.sort(key=lambda ol: -(ol["box_w"] * ol["box_l"]))
        columns = []  # each: {"w":, "l":, "used_height": float, "base": stack dict}
        for ol in cust_lines:
            t = ol["t"]
            remaining = ol["qty"]
            lo, hi = sorted((ol["box_w"], ol["box_l"]))

            for col in (columns if mode == "greedy" else []):
                if remaining <= 0:
                    break
                b_lo, b_hi = sorted((col["w"], col["l"]))
                if not (lo <= b_lo and hi <= b_hi):
                    continue
                n = min(remaining, math.floor((max_total_height - col["used_height"]) / t))
                if n <= 0:
                    continue
                child = {
                    "order_no": ol["order_no"], "cust_code": ol["cust_code"], "cust_name": ol["cust_name"],
                    "lat": ol["lat"], "lng": ol["lng"], "route": ol["route"], "main_road": ol["main_road"],
                    "item_code": ol["item_code"], "delivery_date": ol.get("delivery_date"), "boxes": n, "t": t,
                    "box_w": ol["box_w"], "box_l": ol["box_l"],
                    "height": n * t, "y0": col["used_height"],
                }
                col["base"].setdefault("_children", []).append(child)
                col["base"]["height"] += n * t
                col["used_height"] += n * t
                remaining -= n

            boxes_per_stack = max(1, math.floor(max_total_height / t))
            while remaining > 0:
                n = min(boxes_per_stack, remaining)
                stack = {
                    "order_no": ol["order_no"], "cust_code": ol["cust_code"], "cust_name": ol["cust_name"],
                    "lat": ol["lat"], "lng": ol["lng"], "route": ol["route"], "main_road": ol["main_road"],
                    "item_code": ol["item_code"], "delivery_date": ol.get("delivery_date"), "boxes": n, "t": t,
                    "box_w": ol["box_w"], "box_l": ol["box_l"],
                    "w": ol["box_w"], "l": ol["box_l"],
                    "height": n * t, "y0": 0.0, "_own_height": n * t,
                }
                columns.append({"w": ol["box_w"], "l": ol["box_l"], "used_height": n * t, "base": stack})
                stacks.append(stack)
                remaining -= n
    return stacks


def unmerge_to_order_lines(stacks):
    """Inverse of pack_manual_columns: reconstructs UNCHUNKED order-lines (one per distinct
    order+item+dims, qty = every chunk's boxes summed back together) from a batch of ALREADY-
    merged stacks -- whatever base+rider structure a previous pack_manual_columns run gave them
    is discarded, not read. Used by build()'s per-customer placement loop to re-derive a
    customer's true order quantities so an ALTERNATE stacking strategy can be tried against the
    exact same cargo, instead of being stuck with whichever stacking decision the very first
    (global) pass already committed to."""
    agg = {}

    def add_chunk(s):
        key = (s["cust_code"], s["order_no"], s["item_code"], s["box_w"], s["box_l"], s["t"])
        if key not in agg:
            agg[key] = {
                "order_no": s["order_no"], "cust_code": s["cust_code"], "cust_name": s["cust_name"],
                "lat": s["lat"], "lng": s["lng"], "route": s["route"], "main_road": s["main_road"],
                "item_code": s["item_code"], "delivery_date": s.get("delivery_date"),
                "box_w": s["box_w"], "box_l": s["box_l"], "t": s["t"],
                "qty": 0,
            }
        agg[key]["qty"] += s["boxes"]

    for s in stacks:
        add_chunk(s)
        for c in s.get("_children", []):
            add_chunk(c)
    return list(agg.values())


def expand_items_with_children(items):
    """Unpacks each merged manual-mode column (see pack_manual_columns) back into one report
    row per physical order: the base keeps its own true height, and each rider becomes its own
    entry inheriting the base's floor position (x/y/dx/dy) -- they occupy the same footprint
    slot, just a different height band (y0..y0+height). No-op for stacks that were never
    merged (pallet mode, or a manual stack that never got a rider)."""
    expanded = []
    for it in items:
        children = it.get("_children")
        if not children:
            base = dict(it)
            base.setdefault("y0", 0.0)
            expanded.append(base)
            continue
        base = dict(it)
        base["height"] = base.pop("_own_height", base["height"])
        base.pop("_children", None)
        base.setdefault("y0", 0.0)
        expanded.append(base)
        for child in children:
            expanded.append({**child, "x": it["x"], "y": it["y"], "dx": it["dx"], "dy": it["dy"]})
    return expanded


def _safe_float(value):
    """float(value), or None if blank/invalid -- ERP extracts occasionally have empty
    width_mm/length_mm/thickness_mm/qty_box cells, and those rows should be skipped like
    any other bad-dimension row rather than crashing the whole report."""
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _to_iso_date(value):
    """Informix renders DATE columns through this ODBC/PowerShell path as US-format
    "MM/DD/YYYY HH:MM:SS" (confirmed against a real extract), not ISO -- the HTML report's
    client-side date-range filter does plain string comparison against "YYYY-MM-DD", so passing
    the raw value through unconverted silently matches nothing and the report loads empty. Falls
    back to passing the value through unchanged if it doesn't match the expected shape, rather
    than raising, since a format Informix happens to return differently on some other setup is
    a display quirk, not a reason to crash the whole run."""
    if not value:
        return value
    for fmt in ("%m/%d/%Y %H:%M:%S", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(value.strip(), fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return value


def rows_to_raw_lines(rows):
    """Converts extract_orders.ps1's CSV rows into the compact per-line dicts the HTML report's
    client-side packing engine expects (the same shape the demo Artifact embeds). Shared between
    build()'s static HTML export and server.py's live /api/orders endpoint so the two don't drift
    out of sync with two separate implementations of the same field mapping."""
    raw_lines = []
    for r in rows:
        qty = _safe_float(r["qty_box"])
        w = _safe_float(r["width_mm"]); l = _safe_float(r["length_mm"]); t = _safe_float(r["thickness_mm"])
        if not w or w <= 0 or not l or l <= 0 or not t or t <= 0 or not qty or qty <= 0:
            continue
        order_date_iso = _to_iso_date(r["order_date"])
        delivery_date_iso = _to_iso_date(r.get("delivery_date")) if (r.get("delivery_date") or "").strip() else None
        lat = (r.get("ship_lat") or "").strip()
        lng = (r.get("ship_lng") or "").strip()
        raw_lines.append({
            "o": r["order_no"], "od": order_date_iso, "dd": delivery_date_iso,
            "cc": r["ship_cust_code"].strip(), "cn": r["ship_cust_name"].strip(),
            "lat": float(lat) if lat else None, "lng": float(lng) if lng else None,
            "rt": (r.get("route") or "").strip() or None, "mr": (r.get("main_road") or "").strip() or None,
            "ic": r["item_code"], "q": qty, "w": w, "l": l, "t": t,
            "fl": "", "uc": None,
        })
    return raw_lines


def haversine_km(p1, p2):
    R = 6371.0
    lat1, lng1 = map(math.radians, p1)
    lat2, lng2 = map(math.radians, p2)
    dlat = lat2 - lat1
    dlng = lng2 - lng1
    a = math.sin(dlat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlng / 2) ** 2
    return 2 * R * math.asin(math.sqrt(a))


def _osrm_table_call(points_subset, base_url, timeout):
    coord_str = ";".join(f"{lng},{lat}" for lat, lng in points_subset)
    url = f"{base_url}/table/v1/driving/{coord_str}?annotations=distance"
    with urllib.request.urlopen(url, timeout=timeout) as resp:
        data = json.load(resp)
    if data.get("code") != "Ok":
        raise ValueError(data.get("code", "unknown OSRM error"))
    return data["distances"]


def osrm_distance_matrix_km(codes, points, base_url="http://router.project-osrm.org", timeout=20, chunk=45):
    """Real driving-distance matrix (km) for parallel lists of customer codes / (lat,lng)
    points, via the public OSRM demo server. Returns {code_a: {code_b: km}}, possibly partial,
    or None if every request failed outright -- callers should fall back to haversine_km for
    any missing pair rather than fail the whole report over a network hiccup.

    The public demo caps a single /table request at roughly 100x100 nodes, so on a busy day
    with more than ~50 delivery customers this tiles the codes into groups of `chunk` and
    fetches one request per unordered group pair (covering that pair's union, which yields
    all four sub-blocks at once) instead of a single all-at-once call.
    """
    n = len(points)
    if n < 2:
        return None
    groups = [list(range(i, min(i + chunk, n))) for i in range(0, n, chunk)]
    result = {a: {} for a in codes}
    any_ok = False
    for gi in range(len(groups)):
        for gj in range(gi, len(groups)):
            union = groups[gi] if gi == gj else groups[gi] + groups[gj]
            try:
                dist = _osrm_table_call([points[i] for i in union], base_url, timeout)
            except Exception as e:
                print(f"WARNING: OSRM road-distance lookup failed for group {gi}x{gj} ({e}); those pairs fall back to straight-line distance.")
                continue
            any_ok = True
            for a2, ai in enumerate(union):
                for b2, bi in enumerate(union):
                    d = dist[a2][b2]
                    if d is not None:
                        result[codes[ai]][codes[bi]] = d / 1000.0
    return result if any_ok else None


def two_opt_improve(codes, road_km_fn):
    """2-opt local search over an open path (no fixed start/end -- a truck's stop order isn't
    anchored to a depot here) to shorten the total consecutive-stop distance. Nearest-neighbour
    alone routinely leaves one long doubling-back leg because it never looks past the very next
    stop; with only a handful of stops per truck, exhaustive 2-opt converges in a few passes."""
    if len(codes) < 3:
        return list(codes)

    def path_dist(order):
        return sum(road_km_fn(order[i], order[i + 1]) for i in range(len(order) - 1))

    best = list(codes)
    best_dist = path_dist(best)
    improved = True
    while improved:
        improved = False
        for i in range(len(best) - 1):
            for j in range(i + 1, len(best)):
                cand = best[:i] + best[i:j + 1][::-1] + best[j + 1:]
                cand_dist = path_dist(cand)
                if cand_dist < best_dist - 1e-9:
                    best, best_dist = cand, cand_dist
                    improved = True
    return best


def write_dispatch_xlsx(out_xlsx, header, rows):
    """Formatted .xlsx twin of the dispatch-sheet CSV -- same columns/rows, but with a real
    number type on the numeric columns (so Excel won't mangle order numbers as dates/scientific
    notation the way it sometimes does with plain CSV), bold header, borders, autosized columns,
    a frozen header row and an autofilter so it's usable straight out of the download."""
    numeric_cols = {7, 8, 9, 10, 11}  # 1-indexed: 箱數, 寬mm, 長mm, 厚mm, 堆疊數

    wb = Workbook()
    ws = wb.active
    ws.title = "排車明細"

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(header)
    for col_idx in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in rows:
        ws.append(row)

    max_len = [len(str(h)) for h in header]
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = border
            if c_idx in numeric_cols:
                cell.alignment = Alignment(horizontal="right")
            max_len[c_idx - 1] = max(max_len[c_idx - 1], len(str(value)))

    for c_idx, width in enumerate(max_len, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = min(max(width + 3, 8), 40)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"

    os.makedirs(os.path.dirname(out_xlsx) or ".", exist_ok=True)
    wb.save(out_xlsx)


def build(csv_path, date_label, truck_l, truck_w, truck_h, out_html, out_csv, template_path, out_xlsx=None,
          load_mode="pallet", overhang_w=50.0, overhang_l=300.0, plant_code="T2"):
    with open(csv_path, encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))

    # Build stack units: one physical "stack" = boxes piled to (near) truck height, footprint w x l.
    # Some items (e.g. 平板/紙板 flat paperboard sold by the pallet, no ERP model) are wider than
    # the truck in BOTH horizontal orientations and can never lie flat -- for those, stand them up
    # on edge instead: multiple panels side by side, each consuming its thickness along the floor
    # (bounded by truck_w), with one panel's own edge as height (still bounded by truck_h, same as
    # flat stacking) and the other edge as floor depth (checked against truck_l at placement time,
    # same as flat mode). If neither edge fits within truck_h even standing up, it genuinely can't
    # be loaded in this truck configuration -- skipped (not crashed) and reported at the end.
    #
    # load_mode == "pallet": footprint snaps to a standard pallet SKU plus overhang allowance,
    # base_height reserves room under the cargo for the pallet itself, cargo height capped at
    # MAX_TOTAL_HEIGHT - PALLET_BASE_HEIGHT. load_mode == "manual": no pallet, footprint is just
    # the box's own dimensions, no base height, cargo capped at the full MAX_TOTAL_HEIGHT --
    # matches the demo Artifact's two modes exactly (same constants, same formulas).
    # manual mode's flat (non-vertical) items are deferred to pack_manual_columns below instead
    # of being chunked into stacks right here -- how much of an order's quantity ends up riding
    # on an existing column vs. starting a fresh one isn't known until that pass runs, since it
    # depends on what headroom OTHER same-customer orders happen to leave behind. Vertical
    # (oversized-cross-section) stacks don't participate: their "boxes_per_stack" already means
    # something different (how many panels fit side by side along truck_w, not a height budget
    # each panel accumulates into), so folding them into the same column-filling logic would mix
    # two incompatible chunking rules -- they're rare enough (~1 in 330 lines on a real day) that
    # keeping them as their own independent stacks, same as pallet mode, is the pragmatic choice.
    manual = load_mode == "manual"
    stacks = []
    manual_flat_lines_by_date = {}
    skipped_oversized = []
    for r in rows:
        qty = _safe_float(r["qty_box"])
        w = _safe_float(r["width_mm"]); l = _safe_float(r["length_mm"]); t = _safe_float(r["thickness_mm"])
        if not w or w <= 0 or not l or l <= 0 or not t or t <= 0 or not qty or qty <= 0:
            continue
        lat = (r.get("ship_lat") or "").strip()
        lng = (r.get("ship_lng") or "").strip()
        # 訂單交期(oea101)-- a truck must never carry orders due out on different dates (see
        # pack() below). Falls back to "(無送貨日)" for blank/unset, matching the demo Artifact's
        # own convention exactly, so a customer's orders with no delivery date set still group
        # together as one bucket instead of each silently becoming its own single-order "date".
        delivery_date = _to_iso_date(r.get("delivery_date")) if (r.get("delivery_date") or "").strip() else "(無送貨日)"

        vertical = w > truck_w and l > truck_w
        if vertical:
            height_candidates = [d for d in (w, l) if d <= truck_h]
            if not height_candidates:
                skipped_oversized.append(r)
                continue
            height_dim = max(height_candidates)
            depth_dim = l if height_dim == w else w
            boxes_per_stack = max(1, math.floor(truck_w / t))

            remaining = qty
            while remaining > 0:
                n = min(boxes_per_stack, remaining)
                stacks.append({
                    "order_no": r["order_no"],
                    "cust_code": r["ship_cust_code"].strip(),
                    "cust_name": r["ship_cust_name"].strip(),
                    "lat": float(lat) if lat else None,
                    "lng": float(lng) if lng else None,
                    "route": (r.get("route") or "").strip() or None,
                    "main_road": (r.get("main_road") or "").strip() or None,
                    "item_code": r["item_code"], "delivery_date": delivery_date,
                    "boxes": n, "t": t, "box_w": w, "box_l": l,
                    "w": n * t, "l": depth_dim, "height": height_dim,
                })
                remaining -= n
            continue

        if manual:
            manual_flat_lines_by_date.setdefault(delivery_date, []).append({
                "order_no": r["order_no"],
                "cust_code": r["ship_cust_code"].strip(),
                "cust_name": r["ship_cust_name"].strip(),
                "lat": float(lat) if lat else None,
                "lng": float(lng) if lng else None,
                "route": (r.get("route") or "").strip() or None,
                "main_road": (r.get("main_road") or "").strip() or None,
                "item_code": r["item_code"], "delivery_date": delivery_date,
                "box_w": w, "box_l": l, "t": t, "qty": qty,
            })
            continue

        pallet = choose_pallet(w, l, overhang_w, overhang_l)
        nom_w, nom_l = pallet if pallet else (w, l)
        foot_w, foot_l = nom_w + overhang_w, nom_l + overhang_l
        cols = columns_per_pallet(w, l, foot_w, foot_l)
        boxes_per_stack = cols * max(1, math.floor((MAX_TOTAL_HEIGHT - PALLET_BASE_HEIGHT) / t))

        remaining = qty
        while remaining > 0:
            n = min(boxes_per_stack, remaining)
            layers_used = math.ceil(n / cols)
            stacks.append({
                "order_no": r["order_no"],
                "cust_code": r["ship_cust_code"].strip(),
                "cust_name": r["ship_cust_name"].strip(),
                "lat": float(lat) if lat else None,
                "lng": float(lng) if lng else None,
                "route": (r.get("route") or "").strip() or None,
                "main_road": (r.get("main_road") or "").strip() or None,
                "item_code": r["item_code"], "delivery_date": delivery_date,
                "boxes": n, "t": t, "box_w": w, "box_l": l,
                "w": foot_w, "l": foot_l, "height": PALLET_BASE_HEIGHT + layers_used * t,
            })
            remaining -= n

    # manual-mode same-customer on-top stacking must never merge orders due out on different
    # dates -- calling pack_manual_columns separately PER delivery-date bucket (instead of once
    # across everything) makes that structurally impossible rather than relying on a check
    # somewhere downstream to catch it.
    for delivery_date, lines in manual_flat_lines_by_date.items():
        stacks.extend(pack_manual_columns(lines, MAX_TOTAL_HEIGHT))

    if skipped_oversized:
        print(f"WARNING: {len(skipped_oversized)} order line(s) exceed the truck's cross-section in every orientation (flat or standing) and were excluded from dispatch planning -- these need manual arrangement:")
        for r in skipped_oversized:
            print(f"  order={r['order_no']} cust={r['ship_cust_name'].strip()} item={r['item_code']} "
                  f"dims={r['width_mm']}x{r['length_mm']}x{r['thickness_mm']}mm qty={r['qty_box']}")

    if not stacks:
        raise SystemExit(f"No usable order lines found in {csv_path}")

    # --- route order customers by geographic proximity (occ732 lat / occ733 lng) ---
    # nearest-neighbour greedy chain: start from the westmost-southmost customer, always
    # hop to the closest unvisited one, so consecutive customers in the pack order are
    # geographically close and tend to land on the same truck(s).
    cust_points = {}
    for s in stacks:
        if s["cust_code"] not in cust_points and s["lat"] is not None and s["lng"] is not None:
            cust_points[s["cust_code"]] = (s["lat"], s["lng"])
    no_geo = sorted({s["cust_code"] for s in stacks if s["cust_code"] not in cust_points})

    # real driving-distance matrix between all delivery customers, via OSRM -- used both to
    # order the route (nearest neighbour by road, not as the crow flies) and to report the
    # actual distance a truck will drive. Falls back to haversine per-pair if OSRM is
    # unreachable, so a network hiccup degrades the estimate rather than breaking the report.
    dist_codes = list(cust_points.keys())
    dist_matrix = osrm_distance_matrix_km(dist_codes, [cust_points[c] for c in dist_codes]) if dist_codes else None

    def road_km(code_a, code_b):
        if dist_matrix and code_a in dist_matrix and dist_matrix[code_a].get(code_b) is not None:
            return dist_matrix[code_a][code_b]
        return haversine_km(cust_points[code_a], cust_points[code_b])

    route_order = {}
    if cust_points:
        remaining_custs = dict(cust_points)
        start = min(remaining_custs, key=lambda c: (remaining_custs[c][0], remaining_custs[c][1]))
        order_seq = [start]
        cur_code = start
        remaining_custs.pop(start)
        while remaining_custs:
            nxt = min(remaining_custs, key=lambda c: road_km(cur_code, c))
            order_seq.append(nxt)
            cur_code = nxt
            remaining_custs.pop(nxt)
        route_order = {c: i for i, c in enumerate(order_seq)}
    for i, c in enumerate(no_geo):
        route_order[c] = len(route_order) + i

    cust_name_by_code = {}
    for s in stacks:
        cust_name_by_code.setdefault(s["cust_code"], s["cust_name"])

    cust_route_by_code = {}
    cust_main_road_by_code = {}
    for s in stacks:
        cust_route_by_code.setdefault(s["cust_code"], s["route"])
        cust_main_road_by_code.setdefault(s["cust_code"], s["main_road"])

    def new_truck():
        return {"rows": [], "cur_x": 0.0, "route": None, "main_road": None}

    def clone_truck(truck):
        return {
            "rows": [dict(r, items=list(r["items"])) for r in truck["rows"]],
            "cur_x": truck["cur_x"],
            "route": truck["route"], "main_road": truck["main_road"],
        }

    def try_place(truck, stack):
        w, l = stack["w"], stack["l"]
        for row in truck["rows"]:
            for (dx, dy) in [(l, w), (w, l)]:
                if dy <= row["free_y"] and dx <= row["depth"]:
                    item = dict(stack)
                    item["x"] = row["x0"]; item["y"] = row["used_y"]; item["dx"] = dx; item["dy"] = dy
                    row["items"].append(item)
                    row["used_y"] += dy
                    row["free_y"] -= dy
                    return True
        for (dx, dy) in [(l, w), (w, l)]:
            if dx <= (truck_l - truck["cur_x"]) and dy <= truck_w:
                row = {"x0": truck["cur_x"], "depth": dx, "used_y": 0.0, "free_y": truck_w, "items": []}
                item = dict(stack)
                item["x"] = row["x0"]; item["y"] = 0.0; item["dx"] = dx; item["dy"] = dy
                row["items"].append(item)
                row["used_y"] += dy
                row["free_y"] -= dy
                truck["rows"].append(row)
                truck["cur_x"] += dx
                return True
        return False

    def truck_cargo_vol(truck):
        return sum(it["dx"] * it["dy"] * it["height"] for row in truck["rows"] for it in row["items"])

    def truck_customers(truck):
        return {it["cust_code"] for row in truck["rows"] for it in row["items"]}

    def truck_proximity(truck, cust_code):
        # nearest existing customer on this truck to the candidate customer, road-km -- 0 if
        # the truck is empty or neither point has geo coords (no preference either way then).
        if cust_code not in cust_points:
            return 0.0
        best = None
        for c in truck_customers(truck):
            if c == cust_code or c not in cust_points:
                continue
            d = road_km(c, cust_code)
            if best is None or d < best:
                best = d
        return best if best is not None else 0.0

    # 路線(occ735)/主幹道(occ734): the ERP's own pre-assigned delivery route is a more precise
    # grouping than plain geographic proximity when a customer actually has one, set by the
    # people who actually plan these routes. Only ~21% of customers have a route code though,
    # so a truck locks to whichever grouping mode its FIRST customer determines: "" here means
    # locked to the route-less/open mode (falls back to the plain proximity-only matching this
    # packer always had), any other value means locked to that specific route. The two
    # groupings never mix in one truck.
    def grouping_compatible(truck, cust_code):
        route = cust_route_by_code.get(cust_code)
        if truck["route"] is None:
            return True  # empty so far -- anyone can start it; locking happens in lock_truck_grouping
        if truck["route"] == "":
            return not route  # open-mode truck: only route-less customers
        return route == truck["route"]

    def lock_truck_grouping(truck, cust_code):
        if truck["route"] is not None:
            return
        route = cust_route_by_code.get(cust_code)
        if route:
            truck["route"] = route
            truck["main_road"] = cust_main_road_by_code.get(cust_code)
        else:
            truck["route"] = ""

    def trucks_grouping_compatible(a, b):
        if a["route"] == "" or b["route"] == "":
            return a["route"] == b["route"]
        if a["route"] == b["route"]:
            return True
        # 路線未成車的再按主幹道相同的排車: different specific routes can still combine if they
        # share the same 主幹道 (highway corridor) -- this only ever fires for trucks that still
        # have physical room for each other (checked below), so a route that already filled a
        # truck on its own never gets merged into anything else regardless.
        return bool(a["main_road"]) and a["main_road"] == b["main_road"]

    def consolidate_trucks(trucks):
        """Best-fit-decreasing (even customer-atomic) is still a single forward pass: by the
        time a small, late-processed customer is placed, an earlier truck can look "full
        enough" to reject them even though, in hindsight, two separate under-filled trucks
        could have been one truck all along. This pass looks for exactly that after the fact --
        any two grouping-compatible trucks where ALL of one truck's stacks fit into the other --
        and merges them, repeating (always trying the least-full-by-volume truck first, since
        it's the best candidate to be fully absorbed) until no more merges are possible."""
        merged = True
        while merged:
            merged = False
            trucks.sort(key=truck_cargo_vol)
            for i, donor in enumerate(trucks):
                donor_stacks = [it for row in donor["rows"] for it in row["items"]]
                for j, host in enumerate(trucks):
                    if i == j or not trucks_grouping_compatible(donor, host):
                        continue
                    trial = clone_truck(host)
                    ok = True
                    for it in donor_stacks:
                        stack = {k: v for k, v in it.items() if k not in ("x", "y", "dx", "dy")}
                        if not try_place(trial, stack):
                            ok = False
                            break
                    if not ok:
                        continue
                    trucks[j] = trial
                    del trucks[i]
                    merged = True
                    break
                if merged:
                    break
        return trucks

    # Best-fit-decreasing bin packing, applied at the CUSTOMER level (never split a customer's
    # order across trucks unless it doesn't fit even a single empty one): each customer's whole
    # stack list is tried against every existing truck, and committed to whichever truck can
    # hold ALL of it and leaves the LEAST unused cargo VOLUME afterward (tied-broken by road
    # distance to that truck's existing customers) -- 裝載率優先, 距離其次. Volume, not floor
    # area, since that's what 裝載率 itself means; the per-placement geometry check (try_place)
    # still has to be floor/row-based -- that's real physical feasibility, not a preference.
    TRUCK_VOL = truck_l * truck_w * truck_h
    VOL_BUCKET = TRUCK_VOL * 0.014  # ~1.4% of truck volume: trucks within the same bucket count as "equally full"

    def place_customer(trucks_in, cust_stacks, cust_code):
        """Try cust_stacks against every grouping-compatible existing truck (best-fit-decreasing
        by leftover volume, road-proximity tiebreak), else a fresh truck, else split across
        multiple dedicated fresh trucks. Returns (new_trucks_list, num_new_trucks_added,
        leftover_bucket_or_None) without mutating trucks_in -- existing truck dicts are only
        ever replaced wholesale (via clone_truck) in the RETURNED list, never mutated in place,
        so trying this against the same trucks_in twice with different cust_stacks candidates
        (see the joint stacking/floor-placement comparison below) never cross-contaminates."""
        trucks_out = list(trucks_in)
        best_idx, best_state, best_bucket, best_prox = -1, None, None, None
        for i, truck in enumerate(trucks_out):
            if not grouping_compatible(truck, cust_code):
                continue
            trial = clone_truck(truck)
            ok = True
            for s in cust_stacks:
                if not try_place(trial, s):
                    ok = False
                    break
            if not ok:
                continue
            remaining = TRUCK_VOL - truck_cargo_vol(trial)
            bucket = math.floor(remaining / VOL_BUCKET)
            prox = truck_proximity(truck, cust_code)
            if best_idx < 0 or bucket < best_bucket or (bucket == best_bucket and prox < best_prox):
                best_idx, best_state, best_bucket, best_prox = i, trial, bucket, prox
        if best_idx >= 0:
            trucks_out[best_idx] = best_state
            lock_truck_grouping(trucks_out[best_idx], cust_code)
            return trucks_out, 0, best_bucket

        # no existing truck can take the whole customer -- try a brand new (empty) truck
        # before ever resorting to splitting.
        fresh = new_truck()
        fresh_ok = True
        for s in cust_stacks:
            if not try_place(fresh, s):
                fresh_ok = False
                break
        if fresh_ok:
            trucks_out.append(fresh)
            lock_truck_grouping(trucks_out[-1], cust_code)
            return trucks_out, 1, None

        # doesn't fit any single truck even empty -- this customer's own order genuinely
        # needs more than one truck; fill dedicated fresh trucks one stack at a time.
        added = 0
        cur = new_truck()
        for s in cust_stacks:
            if not try_place(cur, s):
                trucks_out.append(cur)
                lock_truck_grouping(trucks_out[-1], cust_code)
                added += 1
                cur = new_truck()
                if not try_place(cur, s):
                    # genuinely doesn't fit an empty truck in any orientation -- the earlier
                    # flat/vertical pre-check should catch this already, but skip rather than
                    # crash the whole day's run if some other combination slips through.
                    print(f"WARNING: stack does not fit an empty truck in any orientation, excluded: "
                          f"order={s['order_no']} item={s['item_code']} w={s['w']} l={s['l']} height={s['height']}")
                    continue
        trucks_out.append(cur)
        lock_truck_grouping(trucks_out[-1], cust_code)
        added += 1
        return trucks_out, added, None

    # never mix delivery dates on one truck -- group first, then run the customer-clustering/
    # floor-packing pipeline separately per date, same as manual-mode stacking above already
    # does. "(無送貨日)" sorts last so real dates read in order first.
    stacks_by_date = {}
    for s in stacks:
        stacks_by_date.setdefault(s["delivery_date"], []).append(s)
    date_keys = sorted(stacks_by_date.keys(), key=lambda d: (d == "(無送貨日)", d))

    truck_out = []
    truck_no = 0
    for delivery_date in date_keys:
        date_stacks = stacks_by_date[delivery_date]

        by_cust = {}
        for s in date_stacks:
            by_cust.setdefault(s["cust_code"], []).append(s)
        cust_groups = list(by_cust.values())
        for g in cust_groups:
            g.sort(key=lambda s: (-(s["w"] * s["l"]), s["order_no"], s["item_code"]))
        cust_groups.sort(key=lambda g: -sum(s["w"] * s["l"] for s in g))

        # 手工疊車 only: the stacking decision (pack_manual_columns, which orders ride on top of
        # which) and the floor-placement decision (try_place's row packing) are coupled -- a
        # stacking choice made with zero visibility into the floor layout can use up an item that
        # would have completed a row perfectly side by side, and floor placement can only work
        # with whatever units stacking already committed to; it has no way to "unstack" something
        # and try again. Truly solving that jointly (optimal placement with conditional height-
        # merging) is NP-hard in general. What IS tractable: generate an alternate stacking
        # DECISION for this customer's own cargo (no height-stacking at all -- see
        # pack_manual_columns mode="none"), run it through the exact same placement logic as the
        # baseline, and keep whichever candidate's ACTUAL resulting truck count is best (fewer
        # new trucks wins outright; a tie falls back to whichever leaves a fuller existing truck,
        # i.e. the smaller leftover-volume bucket). The baseline (whatever the earlier
        # pack_manual_columns pass already decided) is always one of the candidates, so this can
        # never do worse than not trying at all.
        trucks = []
        for cust_stacks in cust_groups:
            cust_code = cust_stacks[0]["cust_code"]
            candidates = [cust_stacks]
            if manual:
                # vertical (oversized-cross-section) stacks never went through pack_manual_columns
                # in the first place -- their "w"/"l" is a transformed stand-up-on-edge footprint,
                # not the true box_w/box_l unmerge_to_order_lines would reconstruct from, so
                # feeding them back through it would corrupt their packing footprint. Identify
                # them the same way build() originally did (both true dims exceed truck_w) and
                # leave them out of the alternate candidate entirely, unchanged.
                verticals = [s for s in cust_stacks if s["box_w"] > truck_w and s["box_l"] > truck_w]
                flat_stacks = [s for s in cust_stacks if not (s["box_w"] > truck_w and s["box_l"] > truck_w)]
                if flat_stacks:
                    order_lines = unmerge_to_order_lines(flat_stacks)
                    candidates.append(verticals + pack_manual_columns(order_lines, MAX_TOTAL_HEIGHT, mode="none"))

            best_trucks, best_added, best_bucket = None, None, None
            for cand in candidates:
                cand_trucks, added, bucket = place_customer(trucks, cand, cust_code)
                better = (
                    best_added is None
                    or added < best_added
                    or (added == best_added and bucket is not None and (best_bucket is None or bucket < best_bucket))
                )
                if better:
                    best_trucks, best_added, best_bucket = cand_trucks, added, bucket
            trucks = best_trucks

        trucks = consolidate_trucks(trucks)

        for t in trucks:
            truck_no += 1
            # manual mode's merged columns (see pack_manual_columns) travel as one placeable unit
            # through packing/consolidation -- unpack each rider back into its own report row here.
            items = expand_items_with_children([it for row in t["rows"] for it in row["items"]])
            # 裝載率 = 車廂空間使用率(體積),不是樓面使用率(面積) -- 每疊貨的實際體積除以整台
            # 車廂容積,才反映得出這台車實際裝了多少東西,而不是只看地板佔了多少。
            cargo_vol = sum(it["dx"] * it["dy"] * it["height"] for it in items)
            cust_counts = {}
            for it in items:
                key = (it["cust_code"], it["cust_name"])
                cust_counts[key] = cust_counts.get(key, 0) + 1
            primary_cust = max(cust_counts.items(), key=lambda kv: kv[1])[0] if cust_counts else ("", "")
            truck_out.append({
                "truck_no": truck_no, "delivery_date": delivery_date,
                "load_pct": round(cargo_vol / (truck_l * truck_w * truck_h) * 100, 1),
                "num_stacks": len(items),
                "num_customers": len(cust_counts),
                "primary_cust_code": primary_cust[0],
                "primary_cust_name": primary_cust[1],
                "items": items,
            })

    print(f"Total trucks needed: {len(truck_out)}")
    for t in truck_out:
        tag = t["primary_cust_name"] + ("" if t["num_customers"] == 1 else f"+{t['num_customers'] - 1}")
        print(f"  Truck {t['truck_no']:2d} [{t['delivery_date']}]: {t['num_stacks']:3d} stacks, load {t['load_pct']}%  [{tag}]")

    # dispatch sheet (排車單): per truck x order aggregation, grouped by delivery customer
    sheet_header = ["車次", "送貨交期", "送貨客戶代號", "送貨客戶名稱", "訂單號", "品號", "箱數", "寬mm", "長mm", "厚mm", "堆疊數"]
    sheet_rows = []
    for t in truck_out:
        agg = {}
        for it in t["items"]:
            key = (it["cust_code"], it["cust_name"], it["order_no"], it["item_code"], it["box_w"], it["box_l"], it["t"])
            agg.setdefault(key, {"boxes": 0, "stacks": 0})
            agg[key]["boxes"] += it["boxes"]
            agg[key]["stacks"] += 1
        for (cust_code, cust_name, order_no, item, w, l, th), v in sorted(agg.items(), key=lambda kv: route_order[kv[0][0]]):
            sheet_rows.append([f"車次{t['truck_no']}", t["delivery_date"], cust_code, cust_name, order_no, item, int(v["boxes"]), w, l, th, v["stacks"]])

    os.makedirs(os.path.dirname(out_csv) or ".", exist_ok=True)
    with open(out_csv, "w", newline="", encoding="utf-8-sig") as f:
        wr = csv.writer(f)
        wr.writerow(sheet_header)
        wr.writerows(sheet_rows)
    print("Dispatch sheet written:", out_csv)

    if out_xlsx:
        write_dispatch_xlsx(out_xlsx, sheet_header, sheet_rows)
        print("Dispatch sheet (Excel) written:", out_xlsx)

    # --- HTML report: raw order lines + the SAME client-side packing engine the demo Artifact
    # uses, not a pre-packed snapshot. The report used to embed compact_trucks/embed (this
    # module's OWN pack -- try_place, a simple flat shelf packer with no skyline, no multi-
    # strategy, no interactive re-pack) as a static result; the interactive report now computes
    # its own plan in the browser instead (auto-mode AND 模擬裝車 simulation), so what the user
    # actually sees/interacts with is the more capable engine, matching the demo exactly. The
    # CSV/XLSX above are unaffected -- they're still this module's own try_place-based pack,
    # generated independently, so a truck grouping shown in the interactive HTML may not be
    # byte-identical to the CSV row grouping even though both are individually valid.
    raw_lines = rows_to_raw_lines(rows)

    raw_lines_by_plant = {plant_code: {"label": PLANT_NAMES.get(plant_code, plant_code), "lines": raw_lines}}
    order_dates = sorted({rl["od"] for rl in raw_lines if rl["od"]})
    report_defaults = {
        "truckL": truck_l, "truckW": truck_w, "truckH": truck_h,
        "loadMode": load_mode, "plantCode": plant_code, "dateLabel": date_label,
        "rangeStart": order_dates[0] if order_dates else date_label,
        "rangeEnd": order_dates[-1] if order_dates else date_label,
    }

    with open(template_path, encoding="utf-8") as f:
        tpl = f.read()
    html = tpl.replace(
        "/*__RAW_LINES__*/{}/*__END_RAW_LINES__*/",
        json.dumps(raw_lines_by_plant, ensure_ascii=False, separators=(",", ":")),
    )
    html = html.replace(
        '/*__REPORT_DEFAULTS__*/{"truckL":8700,"truckW":2400,"truckH":2400,"loadMode":"pallet","plantCode":"T2","dateLabel":"2026-07-01","rangeStart":"2026-07-01","rangeEnd":"2026-07-01"}/*__END_REPORT_DEFAULTS__*/',
        json.dumps(report_defaults, ensure_ascii=False, separators=(",", ":")),
    )
    html = html.replace("<title>排車單與3D裝載模擬</title>",
                         f"<title>{date_label} 排車與3D裝載模擬 — {PLANT_NAMES.get(plant_code, plant_code)}</title>")

    os.makedirs(os.path.dirname(out_html) or ".", exist_ok=True)
    with open(out_html, "w", encoding="utf-8") as f:
        f.write(html)
    print("HTML report written:", out_html)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--csv", required=True, help="input order-line CSV from extract_orders.ps1")
    ap.add_argument("--date", required=True, help="date label shown in the report, e.g. 2026-07-08")
    ap.add_argument("--template", required=True, help="path to lib/template.html")
    ap.add_argument("--out-html", required=True)
    ap.add_argument("--out-csv", required=True, help="output dispatch-sheet CSV path")
    ap.add_argument("--out-xlsx", help="output dispatch-sheet Excel (.xlsx) path (optional)")
    ap.add_argument("--truck-l", type=float, default=8700.0, help="truck interior length, mm")
    ap.add_argument("--truck-w", type=float, default=2400.0, help="truck interior width, mm")
    ap.add_argument("--truck-h", type=float, default=2400.0, help="truck interior height, mm")
    ap.add_argument("--load-mode", choices=["pallet", "manual"], default="pallet",
                     help="pallet = forklift-loaded on standard pallets (default); manual = hand-stacked, "
                          "no pallet, same-customer orders can stack on top of each other up to 2200mm")
    ap.add_argument("--overhang-w", type=float, default=50.0, help="pallet mode: box overhang allowed past the pallet edge, width axis, mm")
    ap.add_argument("--overhang-l", type=float, default=300.0, help="pallet mode: box overhang allowed past the pallet edge, length axis, mm")
    ap.add_argument("--plant", choices=list(PLANT_NAMES), default="T2",
                     help="plant code (matches config.json's db.database) -- selects the depot "
                          "location the HTML report's client-side route engine measures distance "
                          "from, and which of its four built-in plants the embedded data belongs to")
    args = ap.parse_args()

    build(args.csv, args.date, args.truck_l, args.truck_w, args.truck_h, args.out_html, args.out_csv, args.template, args.out_xlsx,
          load_mode=args.load_mode, overhang_w=args.overhang_w, overhang_l=args.overhang_l, plant_code=args.plant)


if __name__ == "__main__":
    main()
